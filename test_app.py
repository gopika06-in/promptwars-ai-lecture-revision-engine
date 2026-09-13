"""
Test suite for AI Lecture Revision Engine
"""
import io
import sys
import pypdf
import docx
from app import (
    extract_text_from_pdf,
    extract_text_from_docx,
    extract_text_from_txt,
    process_lecture_source,
    parse_sections,
    SUBJECT_PRESETS,
    MAX_CHARACTERS,
)

def test_txt_extraction():
    sample_text = "This is a test lecture on operating systems and memory paging."
    bio = io.BytesIO(sample_text.encode("utf-8"))
    bio.name = "test.txt"
    extracted = extract_text_from_txt(bio)
    assert extracted == sample_text, f"Expected {sample_text}, got {extracted}"
    print("[PASS] TXT extraction test passed")

def test_pdf_extraction():
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    pdf_bytes = io.BytesIO()
    writer.write(pdf_bytes)
    pdf_bytes.seek(0)
    pdf_bytes.name = "sample.pdf"
    
    extracted = extract_text_from_pdf(pdf_bytes)
    print(f"[PASS] PDF extraction test passed (blank page handled smoothly: len={len(extracted)})")

def test_docx_extraction():
    doc = docx.Document()
    doc.add_heading("Biochemistry Lecture 1", level=1)
    doc.add_paragraph("Enzymes catalyze metabolic reactions by lowering activation energy.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Enzyme"
    table.cell(0, 1).text = "Substrate"
    table.cell(1, 0).text = "Amylase"
    table.cell(1, 1).text = "Starch"
    
    docx_bytes = io.BytesIO()
    doc.save(docx_bytes)
    docx_bytes.seek(0)
    docx_bytes.name = "bio.docx"
    
    extracted = extract_text_from_docx(docx_bytes)
    assert "Enzymes catalyze" in extracted
    assert "Amylase | Starch" in extracted
    print("[PASS] DOCX extraction test passed (paragraphs & tables parsed)")

def test_truncation():
    long_text = "A" * 15000
    processed, src, orig_len, is_truncated = process_lecture_source(None, long_text)
    assert orig_len == 15000
    assert is_truncated is True
    assert len(processed) == MAX_CHARACTERS
    print("[PASS] Truncation safety cap test passed (15,000 chars -> 12,000 max)")

def test_section_parsing():
    sample_response = """
# SECTION 1: Core Summary & Executive Key Concepts
- Core summary overview text.
- Key takeaways on machine learning models.

# SECTION 2: Terminology, Key Definitions & Formulas Cheat Sheet
| Term | Definition | Formula |
|---|---|---|
| Loss | Cost function | L(y, y_hat) |

# SECTION 3: 5-Question Multiple Choice Practice Quiz
**Question 1: What is backpropagation?**
A) Forward pass
B) Chain rule gradient calculation
C) Loss function
D) Activation

<details>
<summary>View Answer & Explanation</summary>
**Correct Answer:** B
</details>
"""
    sections = parse_sections(sample_response)
    assert "Core summary overview text" in sections["summary"]
    assert "Loss | Cost function" in sections["cheat_sheet"]
    assert "Question 1" in sections["quiz"]
    print("[PASS] Section parser test passed")

def test_subject_presets():
    assert "General" in SUBJECT_PRESETS
    assert "Bioengineering & Life Sciences" in SUBJECT_PRESETS
    assert "Computer Science & Math" in SUBJECT_PRESETS
    assert "Business & Social Sciences" in SUBJECT_PRESETS
    print("[PASS] Subject presets test passed")

if __name__ == "__main__":
    test_txt_extraction()
    test_pdf_extraction()
    test_docx_extraction()
    test_truncation()
    test_section_parsing()
    test_subject_presets()
    print("\nALL TESTS PASSED SUCCESSFULLY!")
